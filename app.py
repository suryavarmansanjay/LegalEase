import io
import os
import requests
import streamlit as st
from dotenv import load_dotenv
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

load_dotenv()

st.set_page_config(
    page_title="LegalEase",
    page_icon="⚖️",
    layout="wide",
)

BACKEND_URL = os.getenv(
    "BACKEND_URL",
    "http://127.0.0.1:8000"
).rstrip("/")


def create_docx(text: str) -> bytes:
    document = Document()
    document.add_heading("LegalEase - Legal Document Draft", level=1)

    for paragraph in text.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)
        else:
            document.add_paragraph("")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def create_pdf(text: str) -> bytes:
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)

    width, height = A4
    left = 50
    top = height - 50
    y = top

    pdf.setTitle("LegalEase Legal Document Draft")
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(left, y, "LegalEase - Legal Document Draft")
    y -= 30

    pdf.setFont("Helvetica", 10)

    for paragraph in text.split("\n"):
        # Basic wrapping for PDF output.
        words = paragraph.split()
        line = ""

        if not words:
            y -= 14
            continue

        for word in words:
            test_line = f"{line} {word}".strip()
            if pdf.stringWidth(test_line, "Helvetica", 10) <= width - 100:
                line = test_line
            else:
                pdf.drawString(left, y, line)
                y -= 14
                line = word

                if y < 50:
                    pdf.showPage()
                    pdf.setFont("Helvetica", 10)
                    y = top

        if line:
            pdf.drawString(left, y, line)
            y -= 16

        if y < 50:
            pdf.showPage()
            pdf.setFont("Helvetica", 10)
            y = top

    pdf.save()
    return output.getvalue()


st.title("⚖️ LegalEase")
st.subheader("AI Legal Document Generator")

st.info(
    "This application creates AI-generated legal document drafts. "
    "It is not a substitute for legal advice. Review the document "
    "with a qualified legal professional before signing or using it."
)

document_type = st.selectbox(
    "Document Type",
    [
        "Rental Agreement",
        "Employment Agreement",
        "Non-Disclosure Agreement (NDA)",
        "Service Agreement",
        "Sale Agreement",
        "Custom Legal Document",
    ],
)

parties = st.text_area(
    "Parties",
    placeholder="Example: Landlord: John Doe; Tenant: David Kumar",
    height=100,
)

terms = st.text_area(
    "Terms and Conditions",
    placeholder="Enter the important terms, responsibilities, payment details, etc.",
    height=180,
)

dates = st.text_input(
    "Dates",
    placeholder="Example: 01-10-2026 to 30-09-2027",
)

if st.button("⚖️ Generate Document", type="primary"):
    if not parties.strip() or not terms.strip() or not dates.strip():
        st.error("Please complete Parties, Terms and Dates.")
    else:
        payload = {
            "document_type": document_type,
            "parties": parties,
            "terms": terms,
            "dates": dates,
        }

        with st.spinner("Generating your document..."):
            try:
                response = requests.post(
                    f"{BACKEND_URL}/generate",
                    json=payload,
                    timeout=120,
                )

                if response.ok:
                    st.session_state["document"] = response.json()["document"]
                    st.success("Document generated successfully.")
                else:
                    try:
                        detail = response.json().get("detail", response.text)
                    except Exception:
                        detail = response.text
                    st.error(f"Backend error: {detail}")

            except requests.RequestException as exc:
                st.error(
                    "Could not connect to the FastAPI backend. "
                    f"Check BACKEND_URL and make sure the backend is running.\n\n{exc}"
                )

if "document" in st.session_state:
    st.divider()
    st.subheader("Generated Document")

    edited_document = st.text_area(
        "Edit your document before downloading",
        value=st.session_state["document"],
        height=600,
    )

    st.session_state["document"] = edited_document

    txt_data = edited_document.encode("utf-8")
    docx_data = create_docx(edited_document)
    pdf_data = create_pdf(edited_document)

    col1, col2, col3 = st.columns(3)

    with col1:
        st.download_button(
            "⬇️ Download TXT",
            data=txt_data,
            file_name="legal_document.txt",
            mime="text/plain",
            use_container_width=True,
        )

    with col2:
        st.download_button(
            "⬇️ Download DOCX",
            data=docx_data,
            file_name="legal_document.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    with col3:
        st.download_button(
            "⬇️ Download PDF",
            data=pdf_data,
            file_name="legal_document.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
